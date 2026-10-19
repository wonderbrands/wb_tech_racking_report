import odoo
from odoo.api import Environment
from odoo.modules.registry import Registry
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import base64

odoo.tools.config.parse_config(['-c', '/src/odoo.conf', '-d', 'admin'])
registry = Registry('admin')

with registry.cursor() as cr:
    env = Environment(cr, odoo.SUPERUSER_ID, {})
    
    po_names = [
        'PO00054', 'PO00176', 'PO00182', 'PO00166', 'PO00170', 'PO00192', 'PO00193', 'PO00181',
        'PO00110', 'PO00107', 'PO00106', 'P08000', 'PO00147', 'PO00121', 'PO00190', 'PO00055',
        'PO00066', 'PO00184', 'PO00165', 'PO00157', 'PO00172', 'PO00180', 'PO00109', 'PO00167',
        'PO00168', 'PO00105', 'PO00079', 'PO00294', 'PO00070', 'PO00160', 'PO00146', 'PO00159',
        'PO00185', 'PO00186', 'PO00187', 'PO00188', 'PO00163', 'PO00154', 'PO00053'
    ]

    pos = env['purchase.order'].search([('name', 'in', po_names)]).sorted('name')

    po_summary_data = []
    all_skus_data = {}

    for po in pos:
        stor_pickings = env['stock.picking'].sudo().search([
            '|',
            ('purchase_id', '=', po.id),
            ('origin', '=', po.name),
            ('picking_type_id.name', 'ilike', 'Storage')
        ])
        rack_pickings = env['stock.picking'].sudo().search([
            ('origin', '=', po.name),
            ('picking_type_id.name', 'in', ['Rackeo', 'Rackeos', 'Storage'])
        ])
        all_stors = (stor_pickings | rack_pickings).sorted('id')

        move_lines = env['stock.move.line'].sudo().search([
            ('picking_id', 'in', all_stors.ids)
        ])

        po_products = {}
        for line in po.order_line.filtered(lambda l: not l.display_type and l.product_id):
            prod = line.product_id
            sku = prod.default_code or prod.name
            if prod.id not in po_products:
                po_products[prod.id] = {
                    'product': prod,
                    'sku': sku,
                    'ordered_qty': 0.0,
                    'received_qty': 0.0,
                    'racked_qty': 0.0,
                    'waiting_qty': 0.0,
                    'lines': []
                }
            po_products[prod.id]['ordered_qty'] += line.product_qty
            po_products[prod.id]['received_qty'] += line.qty_received

        processed_ml_pickings = set()

        for ml in move_lines:
            prod = ml.product_id
            if not prod:
                continue
            if prod.id not in po_products:
                sku = prod.default_code or prod.name
                po_products[prod.id] = {
                    'product': prod,
                    'sku': sku,
                    'ordered_qty': 0.0,
                    'received_qty': 0.0,
                    'racked_qty': 0.0,
                    'waiting_qty': 0.0,
                    'lines': []
                }

            picking = ml.picking_id
            processed_ml_pickings.add(picking.id)
            user_name = picking.operator.name or picking.create_uid.name or 'Sistema'
            is_automated = picking.validated_by_automation or ('automatizad' in (picking.wmds_log.mapped('log') and ' '.join(picking.wmds_log.mapped('log')).lower() or ''))
            is_done = (picking.state == 'done')
            is_waiting = (picking.state not in ('done', 'cancel'))

            qty = ml.quantity or ml.qty_done or ml.product_uom_qty or 0.0
            move_date = ml.write_date or ml.create_date or picking.date_done or picking.create_date
            date_str = move_date.strftime('%Y-%m-%d %H:%M:%S') if move_date else ''

            line_data = {
                'po': po.name,
                'sku': po_products[prod.id]['sku'],
                'product_name': prod.name,
                'origin_location': ml.location_id.complete_name or ml.location_id.name,
                'dest_location': ml.location_dest_id.complete_name or ml.location_dest_id.name,
                'quantity': qty,
                'ok': '✅' if is_done else ('⏳ En Espera' if is_waiting else '❌ Cancelado'),
                'stor': picking.name,
                'state': 'Realizado' if is_done else ('En Espera' if is_waiting else picking.state),
                'user': user_name,
                'is_automated': 'SI' if is_automated else 'NO',
                'date': date_str,
                'is_done': is_done,
                'is_waiting': is_waiting,
            }

            po_products[prod.id]['lines'].append(line_data)
            if is_done:
                po_products[prod.id]['racked_qty'] += qty
            elif is_waiting:
                po_products[prod.id]['waiting_qty'] += qty

        unprocessed_stors = all_stors.filtered(lambda p: p.id not in processed_ml_pickings and p.state not in ('done', 'cancel'))
        for picking in unprocessed_stors:
            for move in picking.move_ids:
                prod = move.product_id
                if not prod:
                    continue
                if prod.id not in po_products:
                    sku = prod.default_code or prod.name
                    po_products[prod.id] = {
                        'product': prod,
                        'sku': sku,
                        'ordered_qty': 0.0,
                        'received_qty': 0.0,
                        'racked_qty': 0.0,
                        'waiting_qty': 0.0,
                        'lines': []
                    }
                user_name = picking.operator.name or picking.create_uid.name or 'Sistema'
                is_automated = picking.validated_by_automation or ('automatizad' in (picking.wmds_log.mapped('log') and ' '.join(picking.wmds_log.mapped('log')).lower() or ''))
                qty = move.product_uom_qty or move.quantity or 0.0
                move_date = move.write_date or move.create_date or picking.create_date
                date_str = move_date.strftime('%Y-%m-%d %H:%M:%S') if move_date else ''

                line_data = {
                    'po': po.name,
                    'sku': po_products[prod.id]['sku'],
                    'product_name': prod.name,
                    'origin_location': move.location_id.complete_name or move.location_id.name,
                    'dest_location': move.location_dest_id.complete_name or move.location_dest_id.name,
                    'quantity': qty,
                    'ok': '⏳ En Espera',
                    'stor': picking.name,
                    'state': 'En Espera',
                    'user': user_name,
                    'is_automated': 'SI' if is_automated else 'NO',
                    'date': date_str,
                    'is_done': False,
                    'is_waiting': True,
                }
                po_products[prod.id]['lines'].append(line_data)
                po_products[prod.id]['waiting_qty'] += qty

        po_has_issues = False
        po_tot_ordered = 0.0
        po_tot_received = 0.0
        po_tot_racked = 0.0
        po_tot_waiting = 0.0
        po_tot_missing = 0.0
        po_tot_overdemand = 0.0

        for prod_id, pdata in po_products.items():
            racked = pdata['racked_qty']
            waiting = pdata['waiting_qty']
            ordered = pdata['ordered_qty']
            received = pdata['received_qty']
            base_target = received if received > 0 else ordered
            missing = max(0.0, base_target - racked)
            overdemand = max(0.0, (racked + waiting) - base_target)

            po_tot_ordered += ordered
            po_tot_received += received
            po_tot_racked += racked
            po_tot_waiting += waiting
            po_tot_missing += missing
            po_tot_overdemand += overdemand

            if missing > 0 or overdemand > 0 or waiting > 0:
                po_has_issues = True

                sku_key = pdata['sku']
                if sku_key not in all_skus_data:
                    all_skus_data[sku_key] = {
                        'product_name': pdata['product'].name,
                        'sku': sku_key,
                        'ordered_qty': 0.0,
                        'received_qty': 0.0,
                        'racked_qty': 0.0,
                        'waiting_qty': 0.0,
                        'missing_qty': 0.0,
                        'overdemand_qty': 0.0,
                        'lines': []
                    }
                all_skus_data[sku_key]['ordered_qty'] += ordered
                all_skus_data[sku_key]['received_qty'] += received
                all_skus_data[sku_key]['racked_qty'] += racked
                all_skus_data[sku_key]['waiting_qty'] += waiting
                all_skus_data[sku_key]['missing_qty'] += missing
                all_skus_data[sku_key]['overdemand_qty'] += overdemand
                all_skus_data[sku_key]['lines'].extend(pdata['lines'])

        if po_has_issues:
            if po_tot_overdemand > 0:
                status = f"sobre demanda de rackeo en {int(po_tot_overdemand) if po_tot_overdemand.is_integer() else po_tot_overdemand}, borrar"
            elif po_tot_missing > 0:
                if po_tot_waiting > 0:
                    status = f"⚠️ Incompleto / Faltan {int(po_tot_missing) if po_tot_missing.is_integer() else po_tot_missing} pzs ({int(po_tot_waiting) if po_tot_waiting.is_integer() else po_tot_waiting} en espera)"
                else:
                    status = f"⚠️ Incompleto / Faltan {int(po_tot_missing) if po_tot_missing.is_integer() else po_tot_missing} pzs"
            else:
                status = '✅ Acomodado Completo'

            po_summary_data.append({
                'po': po.name,
                'partner': po.partner_id.name or '',
                'state': po.state,
                'ordered': po_tot_ordered,
                'received': po_tot_received,
                'racked': po_tot_racked,
                'waiting': po_tot_waiting,
                'missing': po_tot_missing,
                'overdemand': po_tot_overdemand,
                'status': status
            })

    wb = openpyxl.Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1F497D', end_color='1F497D', fill_type='solid')
    sub_font = Font(name='Calibri', size=11, bold=True)
    over_demand_font = Font(name='Calibri', size=11, bold=True, color='C00000')
    waiting_font = Font(name='Calibri', size=11, bold=True, color='E67E22')
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')

    # SHEET 1: Resumen_Consolidado_POs
    ws_summary = wb.create_sheet(title='Resumen_Consolidado_POs')
    ws_summary.append(['REPORTE CONSOLIDADO DE ORDENES DE COMPRA INCOMPLETAS / SOBREDEMANDAS'])
    ws_summary.append([f'Total POs Analizadas: {len(pos)}', f'POs con Inconsistencias: {len(po_summary_data)}'])
    ws_summary.append([])

    summary_headers = ['PO', 'Proveedor', 'Estado PO', 'Cant. Ordenada', 'Cant. Recibida', 'Rackeado (Realizado)', 'En Espera (Pendiente)', 'Faltantes', 'Sobredemanda', 'Estado Rackeo']
    ws_summary.append(summary_headers)

    for col_num in range(1, len(summary_headers) + 1):
        cell = ws_summary.cell(row=4, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center

    tot_ord = tot_rec = tot_rac = tot_wai = tot_mis = tot_ove = 0.0

    curr_row = 5
    for pdata in po_summary_data:
        ws_summary.append([
            pdata['po'],
            pdata['partner'],
            pdata['state'],
            pdata['ordered'],
            pdata['received'],
            pdata['racked'],
            pdata['waiting'],
            pdata['missing'],
            pdata['overdemand'],
            pdata['status']
        ])

        ws_summary.cell(row=curr_row, column=1).alignment = align_left
        ws_summary.cell(row=curr_row, column=2).alignment = align_left
        ws_summary.cell(row=curr_row, column=3).alignment = align_center
        ws_summary.cell(row=curr_row, column=4).alignment = align_right
        ws_summary.cell(row=curr_row, column=5).alignment = align_right
        ws_summary.cell(row=curr_row, column=6).alignment = align_right
        ws_summary.cell(row=curr_row, column=7).alignment = align_right
        ws_summary.cell(row=curr_row, column=8).alignment = align_right
        ws_summary.cell(row=curr_row, column=9).alignment = align_right
        ws_summary.cell(row=curr_row, column=10).alignment = align_center

        if "borrar" in pdata['status']:
            ws_summary.cell(row=curr_row, column=10).font = over_demand_font
        elif "Incompleto" in pdata['status']:
            ws_summary.cell(row=curr_row, column=10).font = waiting_font

        for c in range(1, 11):
            ws_summary.cell(row=curr_row, column=c).border = thin_border

        tot_ord += pdata['ordered']
        tot_rec += pdata['received']
        tot_rac += pdata['racked']
        tot_wai += pdata['waiting']
        tot_mis += pdata['missing']
        tot_ove += pdata['overdemand']
        curr_row += 1

    # Totals row
    ws_summary.append([
        'TOTALES',
        '',
        '',
        tot_ord,
        tot_rec,
        tot_rac,
        tot_wai,
        tot_mis,
        tot_ove,
        '⚠️ REVISAR POs CON PENDIENTES / SOBREDEMANDAS' if (tot_mis > 0 or tot_ove > 0) else '✅ TODO COMPLETO'
    ])

    for c in range(1, 11):
        cell = ws_summary.cell(row=curr_row, column=c)
        cell.font = sub_font
        cell.border = thin_border

    # INDIVIDUAL SKU SHEETS
    used_sheet_names = set(['Resumen_Consolidado_POs'])

    for sku_key, sdata in all_skus_data.items():
        clean_sku = "".join([c for c in sku_key if c.isalnum() or c in ('-', '_')]).strip()[:28]
        sheet_title = clean_sku or f'SKU_INC'

        base_title = sheet_title
        counter = 1
        while sheet_title in used_sheet_names:
            sheet_title = f"{base_title[:25]}_{counter}"
            counter += 1
        used_sheet_names.add(sheet_title)

        ws_prod = wb.create_sheet(title=sheet_title)

        racked = sdata['racked_qty']
        waiting = sdata['waiting_qty']
        ordered = sdata['ordered_qty']
        received = sdata['received_qty']
        base_target = received if received > 0 else ordered
        missing = sdata['missing_qty']
        overdemand = sdata['overdemand_qty']

        ws_prod.append([f"DESGLOSE DE MOVIMIENTOS RACKEO - SKU: {sdata['sku']}"])
        ws_prod.append([f"Producto: {sdata['product_name']}"])

        if overdemand > 0:
            status_header_str = f"SOBREDEMANDA: sobre demanda de rackeo en {int(overdemand) if overdemand.is_integer() else overdemand}, borrar"
        elif missing > 0:
            if waiting > 0:
                status_header_str = f"ESTADO: Faltan {int(missing) if missing.is_integer() else missing} pzs por rackear ({int(waiting) if waiting.is_integer() else waiting} en espera)"
            else:
                status_header_str = f"ESTADO: Faltan {int(missing) if missing.is_integer() else missing} pzs por rackear"
        else:
            status_header_str = "ESTADO: ✅ Acomodado Completo"

        ws_prod.append([f"Ordenado: {ordered}", f"Recibido: {received}", f"Rackeado: {racked}", f"En Espera: {waiting}", status_header_str])
        ws_prod.append([])

        prod_headers = ['PO', 'SKU', 'FECHA Y HORA', 'ORIGEN', 'DESTINO (UBICACIÓN)', 'PZS', 'Ok / Estado', 'STOR', 'Estado Picking', 'Usuario', 'Automatizado', 'Observación']
        ws_prod.append(prod_headers)

        for col_num in range(1, len(prod_headers) + 1):
            cell = ws_prod.cell(row=5, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center

        p_row = 6
        if sdata['lines']:
            accum_total = 0.0
            for line in sdata['lines']:
                line_qty = line['quantity']
                accum_total += line_qty

                if accum_total > base_target and base_target > 0:
                    excess = min(line_qty, accum_total - base_target)
                    obs = f"sobre demanda de rackeo en {int(excess) if excess.is_integer() else excess}, borrar"
                elif line.get('is_waiting'):
                    obs = "en espera"
                else:
                    obs = "totales"

                ws_prod.append([
                    line['po'],
                    line['sku'],
                    line['date'],
                    line['origin_location'],
                    line['dest_location'],
                    line['quantity'],
                    line['ok'],
                    line['stor'],
                    line['state'],
                    line['user'],
                    line['is_automated'],
                    obs
                ])

                ws_prod.cell(row=p_row, column=1).alignment = align_left
                ws_prod.cell(row=p_row, column=2).alignment = align_left
                ws_prod.cell(row=p_row, column=3).alignment = align_center
                ws_prod.cell(row=p_row, column=4).alignment = align_left
                ws_prod.cell(row=p_row, column=5).alignment = align_left
                ws_prod.cell(row=p_row, column=6).alignment = align_right
                ws_prod.cell(row=p_row, column=7).alignment = align_center
                ws_prod.cell(row=p_row, column=8).alignment = align_left
                ws_prod.cell(row=p_row, column=9).alignment = align_center
                ws_prod.cell(row=p_row, column=10).alignment = align_left
                ws_prod.cell(row=p_row, column=11).alignment = align_center
                ws_prod.cell(row=p_row, column=12).alignment = align_left

                if "borrar" in obs:
                    ws_prod.cell(row=p_row, column=12).font = over_demand_font
                elif line.get('is_waiting'):
                    ws_prod.cell(row=p_row, column=7).font = waiting_font

                for c in range(1, 13):
                    ws_prod.cell(row=p_row, column=c).border = thin_border
                p_row += 1
        else:
            ws_prod.append(['- Sin movimientos de rackeo registrados -'])

    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 65)

    output_path = '/extra-addons/wonderbrands2026/Reporte_Rackeo_Consolidado_POs_Incompletas.xlsx'
    wb.save(output_path)
    print(f"EXCEL_SUCCESS: Saved consolidated report to {output_path}")

    # Attach to ir.attachment
    with open(output_path, 'rb') as f:
        file_bytes = f.read()

    att = env['ir.attachment'].sudo().create({
        'name': 'Reporte_Rackeo_Consolidado_POs_Incompletas.xlsx',
        'type': 'binary',
        'datas': base64.b64encode(file_bytes),
        'res_model': 'purchase.order',
        'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    })
    print(f"ATTACHMENT_CREATED: Attachment ID {att.id}")
    env.cr.commit()
    print("CR_COMMITTED_SUCCESSFULLY")
