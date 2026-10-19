import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import io
import base64

po_names = [
    'PO00054', 'PO00176', 'PO00182', 'PO00166', 'PO00170', 'PO00192', 'PO00193', 'PO00181',
    'PO00110', 'PO00107', 'PO00106', 'P08000', 'PO00147', 'PO00121', 'PO00190', 'PO00055',
    'PO00066', 'PO00184', 'PO00165', 'PO00157', 'PO00172', 'PO00180', 'PO00109', 'PO00167',
    'PO00168', 'PO00105', 'PO00079', 'PO00294', 'PO00070', 'PO00160', 'PO00146', 'PO00159',
    'PO00185', 'PO00186', 'PO00187', 'PO00188', 'PO00163', 'PO00154', 'PO00053'
]

pos = env['purchase.order'].search([('name', 'in', po_names)]).sorted('name')

wb = openpyxl.Workbook()
default_sheet = wb.active
wb.remove(default_sheet)

# Estilos Excel
title_font = Font(name='Calibri', size=14, bold=True, color='1F497D')
header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='1F497D', end_color='1F497D', fill_type='solid')
sub_header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
sub_font = Font(name='Calibri', size=11, bold=True)

over_demand_font = Font(name='Calibri', size=11, bold=True, color='C00000')
waiting_font = Font(name='Calibri', size=11, bold=True, color='E67E22')
ok_font = Font(name='Calibri', size=11, bold=True, color='375623')

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)
double_bottom_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='double', color='000000')
)

align_center = Alignment(horizontal='center', vertical='center')
align_left = Alignment(horizontal='left', vertical='center')
align_right = Alignment(horizontal='right', vertical='center')

# --- HOJA 1: RESUMEN DE LINEAS DE SKU CON ERROR ---
ws_summary = wb.create_sheet(title='Resumen_SKUs_Con_Error')
ws_summary.append(['REPORTE CONSOLIDADO DE LÍNEAS DE SKU CON ERRORES E INCONSISTENCIAS'])
ws_summary.cell(row=1, column=1).font = title_font
ws_summary.append([f'Total POs Auditadas: {len(pos)}', 'Filtro: Exclusivamente líneas con inconsistencias en Recepción o Rackeo'])
ws_summary.append([])

summary_headers = [
    'PO', 'SKU', 'Producto', 'Cant. Ordenada (PO)', 'Cant. Recibida (IN)',
    'Dif. Recepción vs PO', 'Estado Recepción',
    'Rackeado (Realizado)', 'En Espera (Pendiente)',
    'Faltantes (vs PO)', 'Sobredemanda (vs PO)',
    'Estado Rackeo (vs PO)', 'Estado Rackeo (vs Recepción)'
]
ws_summary.append(summary_headers)

for col_num in range(1, len(summary_headers) + 1):
    cell = ws_summary.cell(row=4, column=col_num)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = align_center

summary_row_idx = 5
total_ordered_all = 0.0
total_received_all = 0.0
total_racked_all = 0.0
total_waiting_all = 0.0
total_missing_po_all = 0.0
total_overdemand_po_all = 0.0

all_error_sku_sheets = {}

for po in pos:
    stor_pickings = env['stock.picking'].sudo().search([
        '|', ('purchase_id', '=', po.id), ('origin', '=', po.name),
        ('picking_type_id.name', 'ilike', 'Storage')
    ])
    rack_pickings = env['stock.picking'].sudo().search([
        ('origin', '=', po.name),
        ('picking_type_id.name', 'in', ['Rackeo', 'Rackeos', 'Storage'])
    ])
    all_stors = (stor_pickings | rack_pickings).sorted('id')
    move_lines = env['stock.move.line'].sudo().search([('picking_id', 'in', all_stors.ids)])

    prod_to_parent = {}
    po_products = {}

    for line in po.order_line.filtered(lambda l: not l.display_type and l.product_id):
        parent = line.product_id
        sku = parent.default_code or parent.name
        prod_to_parent[parent.id] = parent.id
        
        comp_prods = set([parent.id])
        prefix = (parent.default_code or '').strip()
        
        prefix_boxes = env['product.product'].sudo().search([('default_code', '=like', f"{prefix}#%")]) if prefix else env['product.product']
        if prefix_boxes:
            for box in prefix_boxes:
                comp_prods.add(box.id)
                prod_to_parent[box.id] = parent.id
        else:
            bom = env['mrp.bom'].sudo().search([
                '|', ('product_id', '=', parent.id), ('product_tmpl_id', '=', parent.product_tmpl_id.id),
                ('type', '=', 'phantom')
            ], limit=1)
            if bom and bom.bom_line_ids:
                for bl in bom.bom_line_ids:
                    comp_prods.add(bl.product_id.id)
                    prod_to_parent[bl.product_id.id] = parent.id

        if parent.id not in po_products:
            po_products[parent.id] = {
                'product': parent,
                'sku': sku,
                'ordered_qty': 0.0,
                'received_qty': 0.0,
                'racked_qty': 0.0,
                'waiting_qty': 0.0,
                'comp_ids': comp_prods,
                'comp_racked_done': {cid: 0.0 for cid in comp_prods},
                'comp_racked_waiting': {cid: 0.0 for cid in comp_prods},
                'lines': []
            }
        po_products[parent.id]['ordered_qty'] += line.product_qty
        po_products[parent.id]['received_qty'] += line.qty_received

    for ml in move_lines:
        prod = ml.product_id
        if not prod:
            continue
        parent_id = prod_to_parent.get(prod.id)
        if not parent_id or parent_id not in po_products:
            sku = prod.default_code or prod.name
            po_products[prod.id] = {
                'product': prod,
                'sku': sku,
                'ordered_qty': 0.0,
                'received_qty': 0.0,
                'racked_qty': 0.0,
                'waiting_qty': 0.0,
                'comp_ids': set([prod.id]),
                'comp_racked_done': {prod.id: 0.0},
                'comp_racked_waiting': {prod.id: 0.0},
                'lines': []
            }
            parent_id = prod.id
            prod_to_parent[prod.id] = prod.id
        pdata = po_products[parent_id]
        picking = ml.picking_id
        user_name = picking.operator.name or picking.create_uid.name or 'Sistema'
        is_automated = picking.validated_by_automation or ('automatizad' in (picking.wmds_log.mapped('log') and ' '.join(picking.wmds_log.mapped('log')).lower() or ''))
        is_done = (picking.state == 'done')
        is_waiting = (picking.state not in ('done', 'cancel'))
        qty = ml.quantity or ml.qty_done or ml.product_uom_qty or 0.0
        move_date = ml.write_date or ml.create_date or picking.date_done or picking.create_date
        date_str = move_date.strftime('%Y-%m-%d %H:%M:%S') if move_date else ''

        if prod.id in pdata['comp_racked_done']:
            if is_done:
                pdata['comp_racked_done'][prod.id] += qty
            elif is_waiting:
                pdata['comp_racked_waiting'][prod.id] += qty

        line_data = {
            'po': po.name,
            'parent_sku': pdata['sku'],
            'sku': prod.default_code or prod.name,
            'product_name': prod.name,
            'origin_location': ml.location_id.complete_name or ml.location_id.name,
            'dest_location': ml.location_dest_id.complete_name or ml.location_dest_id.name,
            'quantity': qty,
            'ok': '✅' if is_done else ('⏳ En Espera' if is_waiting else '❌ Cancelado'),
            'stor': picking.name,
            'state': 'Realizado' if is_done else ('En Espera' if is_waiting else picking.state),
            'user': user_name,
            'is_automated': 'SI' if is_automated else 'NO',
            'date': date_str,
            'is_done': is_done,
            'is_waiting': is_waiting,
        }
        pdata['lines'].append(line_data)

    for parent_id, pdata in po_products.items():
        box_cids = [cid for cid in pdata['comp_ids'] if cid != parent_id]
        if box_cids:
            pdata['racked_qty'] = min(pdata['comp_racked_done'][cid] for cid in box_cids)
            pdata['waiting_qty'] = min(pdata['comp_racked_waiting'][cid] for cid in box_cids)
        else:
            pdata['racked_qty'] = pdata['comp_racked_done'].get(parent_id, 0.0)
            pdata['waiting_qty'] = pdata['comp_racked_waiting'].get(parent_id, 0.0)

        ordered = pdata['ordered_qty']
        received = pdata['received_qty']
        racked = pdata['racked_qty']
        waiting = pdata['waiting_qty']
        total_processed = racked + waiting

        base_target_po = ordered if ordered > 0 else received
        missing_po = max(0.0, base_target_po - racked)
        overdemand_po = max(0.0, total_processed - base_target_po)

        base_target_rec = received if received > 0 else ordered
        overdemand_rec = max(0.0, total_processed - base_target_rec)
        missing_rec_rack = max(0.0, base_target_rec - racked)

        diff_rec = received - ordered

        has_error = (overdemand_po > 0) or (missing_po > 0) or (diff_rec != 0 and ordered > 0)

        if has_error:
            if received > ordered and ordered > 0:
                status_rec = f"⚠️ Exceso Recibido (+{int(diff_rec)} pzs vs PO)"
            elif received < ordered and ordered > 0:
                status_rec = f"⚠️ Recepción Incompleta (-{int(ordered - received)} pzs vs PO)"
            else:
                status_rec = "✅ Recepción Completa"

            if overdemand_po > 0:
                status_rack_po = f"sobre demanda de rackeo en {int(overdemand_po)}, borrar"
            elif missing_po > 0:
                if waiting > 0:
                    status_rack_po = f"⚠️ Incompleto / Faltan {int(missing_po)} pzs ({int(waiting)} en espera)"
                else:
                    status_rack_po = f"⚠️ Incompleto / Faltan {int(missing_po)} pzs"
            else:
                status_rack_po = "✅ Acomodado Completo"

            if overdemand_rec > 0:
                status_rack_rec = f"Exceso STOR vs Recepción ({int(overdemand_rec)} pzs)"
            elif missing_rec_rack > 0:
                status_rack_rec = f"Pendiente de Acomodar ({int(missing_rec_rack)} pzs por rackear)"
            else:
                status_rack_rec = "✅ Acomodado Completo (100% de lo recibido)"

            ws_summary.append([
                po.name,
                pdata['sku'],
                pdata['product'].name,
                ordered,
                received,
                diff_rec,
                status_rec,
                racked,
                waiting,
                missing_po,
                overdemand_po,
                status_rack_po,
                status_rack_rec
            ])

            ws_summary.cell(row=summary_row_idx, column=1).alignment = align_center
            ws_summary.cell(row=summary_row_idx, column=2).alignment = align_left
            ws_summary.cell(row=summary_row_idx, column=3).alignment = align_left
            ws_summary.cell(row=summary_row_idx, column=4).alignment = align_right
            ws_summary.cell(row=summary_row_idx, column=5).alignment = align_right
            ws_summary.cell(row=summary_row_idx, column=6).alignment = align_right
            ws_summary.cell(row=summary_row_idx, column=7).alignment = align_center
            ws_summary.cell(row=summary_row_idx, column=8).alignment = align_right
            ws_summary.cell(row=summary_row_idx, column=9).alignment = align_right
            ws_summary.cell(row=summary_row_idx, column=10).alignment = align_right
            ws_summary.cell(row=summary_row_idx, column=11).alignment = align_right
            ws_summary.cell(row=summary_row_idx, column=12).alignment = align_left
            ws_summary.cell(row=summary_row_idx, column=13).alignment = align_left

            if overdemand_po > 0:
                ws_summary.cell(row=summary_row_idx, column=11).font = over_demand_font
                ws_summary.cell(row=summary_row_idx, column=12).font = over_demand_font
            if missing_po > 0:
                ws_summary.cell(row=summary_row_idx, column=10).font = waiting_font
                ws_summary.cell(row=summary_row_idx, column=12).font = waiting_font

            for col_c in range(1, 14):
                ws_summary.cell(row=summary_row_idx, column=col_c).border = thin_border

            total_ordered_all += ordered
            total_received_all += received
            total_racked_all += racked
            total_waiting_all += waiting
            total_missing_po_all += missing_po
            total_overdemand_po_all += overdemand_po
            summary_row_idx += 1

            sku = pdata['sku']
            if sku not in all_error_sku_sheets:
                all_error_sku_sheets[sku] = []
            all_error_sku_sheets[sku].append({
                'po': po.name,
                'pdata': pdata,
                'missing_po': missing_po,
                'overdemand_po': overdemand_po
            })

# Fila de Totales Generales
ws_summary.append([
    'TOTALES SKUS CON ERROR', None, None,
    total_ordered_all, total_received_all,
    total_received_all - total_ordered_all,
    '⚠️ FALTANTES/EXCESOS EN RECEPCIÓN' if total_received_all != total_ordered_all else '✅ RECEPCIÓN IGUAL',
    total_racked_all, total_waiting_all,
    total_missing_po_all, total_overdemand_po_all,
    '⚠️ ERRORES DE RACKEO DETECTADOS',
    '⚠️ REVISAR PENDIENTES Y SOBREDEMANDAS'
])
tot_row = summary_row_idx
for col_c in range(1, 14):
    cell = ws_summary.cell(row=tot_row, column=col_c)
    cell.font = sub_font
    cell.fill = sub_header_fill
    cell.border = double_bottom_border

ws_summary.cell(row=tot_row, column=1).alignment = align_left
ws_summary.cell(row=tot_row, column=4).alignment = align_right
ws_summary.cell(row=tot_row, column=5).alignment = align_right
ws_summary.cell(row=tot_row, column=6).alignment = align_right
ws_summary.cell(row=tot_row, column=8).alignment = align_right
ws_summary.cell(row=tot_row, column=9).alignment = align_right
ws_summary.cell(row=tot_row, column=10).alignment = align_right
ws_summary.cell(row=tot_row, column=11).alignment = align_right
ws_summary.cell(row=tot_row, column=11).font = over_demand_font
ws_summary.cell(row=tot_row, column=10).font = waiting_font

# --- HOJAS POR SKU DE SKUS CON ERROR ---
for sku, items in all_error_sku_sheets.items():
    clean_sku = "".join([c for c in sku if c.isalnum() or c in ('-', '_')]).strip()[:30]
    ws_prod = wb.create_sheet(title=clean_sku)

    ws_prod.append([f'DETALLE DE MOVIMIENTOS DE RACKEO (STOR) - SKU CON ERROR: {sku}'])
    ws_prod.cell(row=1, column=1).font = title_font
    ws_prod.append([])

    prod_headers = [
        'PO', 'SKU', 'FECHA Y HORA', 'ORIGEN', 'DESTINO (UBICACIÓN)',
        'PZS', 'Ok / Estado', 'STOR', 'Estado Picking', 'Usuario', 'Automatizado', 'Observación'
    ]
    ws_prod.append(prod_headers)
    for col_c in range(1, 13):
        cell = ws_prod.cell(row=3, column=col_c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center

    p_row = 4
    for item in items:
        pdata = item['pdata']
        lines = pdata['lines']
        ordered = pdata['ordered_qty']
        racked_accum = 0.0

        for line in lines:
            if line['is_done']:
                racked_accum += line['quantity']
                if ordered > 0 and racked_accum > ordered:
                    excess = racked_accum - ordered
                    obs = f"sobre demanda de rackeo en {int(excess) if excess.is_integer() else excess}, borrar"
                else:
                    obs = "totales"
            elif line.get('is_waiting'):
                obs = "en espera"
            else:
                obs = "totales"

            ws_prod.append([
                line['po'],
                line['sku'],
                line['date'],
                line['origin_location'],
                line['dest_location'],
                line['quantity'],
                line['ok'],
                line['stor'],
                line['state'],
                line['user'],
                line['is_automated'],
                obs
            ])

            ws_prod.cell(row=p_row, column=1).alignment = align_left
            ws_prod.cell(row=p_row, column=2).alignment = align_left
            ws_prod.cell(row=p_row, column=3).alignment = align_center
            ws_prod.cell(row=p_row, column=4).alignment = align_left
            ws_prod.cell(row=p_row, column=5).alignment = align_left
            ws_prod.cell(row=p_row, column=6).alignment = align_right
            ws_prod.cell(row=p_row, column=7).alignment = align_center
            ws_prod.cell(row=p_row, column=8).alignment = align_left
            ws_prod.cell(row=p_row, column=9).alignment = align_center
            ws_prod.cell(row=p_row, column=10).alignment = align_left
            ws_prod.cell(row=p_row, column=11).alignment = align_center
            ws_prod.cell(row=p_row, column=12).alignment = align_left

            if "borrar" in obs:
                ws_prod.cell(row=p_row, column=12).font = over_demand_font
            elif line.get('is_waiting'):
                ws_prod.cell(row=p_row, column=7).font = waiting_font

            for col_c in range(1, 13):
                ws_prod.cell(row=p_row, column=col_c).border = thin_border
            p_row += 1

# Auto-ajustar ancho de columnas en todas las hojas
for ws in wb.worksheets:
    for col in ws.columns:
        max_len = 0
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 65)

output_path = "/extra-addons/Reporte_Rackeo_Consolidado_Solo_Errores.xlsx"
wb.save(output_path)
print(f"EXCEL_ERROR_LINES_SUCCESS: Saved to {output_path}")

# Guardar en ir.attachment
output_bytes = io.BytesIO()
wb.save(output_bytes)
file_data = base64.b64encode(output_bytes.getvalue())

att = env['ir.attachment'].sudo().create({
    'name': 'Reporte_Rackeo_Consolidado_Solo_Errores.xlsx',
    'datas': file_data,
    'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'public': True,
})
print(f"ATTACHMENT_ERROR_LINES_CREATED: Attachment ID {att.id}")
env.cr.commit()
print("CR_COMMITTED_SUCCESSFULLY")
