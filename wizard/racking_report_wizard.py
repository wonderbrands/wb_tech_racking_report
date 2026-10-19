# -*- coding: utf-8 -*-
import io
import base64
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from odoo import models, fields, api, _
from odoo.exceptions import UserError


class WbRackingReportWizard(models.TransientModel):
    _name = 'wb.racking.report.wizard'
    _description = 'Wizard de Reporte de Rackeo Excel'

    purchase_id = fields.Many2one('purchase.order', string='Orden de Compra', required=True)
    report_file = fields.Binary('Archivo Excel', readonly=True)
    file_name = fields.Char('Nombre del Archivo', readonly=True)

    def action_generate_excel_report(self):
        self.ensure_one()
        po = self.purchase_id

        # 1. Buscar todos los pickings de STOR/Rackeo asociados
        stor_pickings = self.env['stock.picking'].sudo().search([
            '|',
            ('purchase_id', '=', po.id),
            ('origin', '=', po.name),
            ('picking_type_id.name', 'ilike', 'Storage')
        ])
        
        rack_pickings = self.env['stock.picking'].sudo().search([
            ('origin', '=', po.name),
            ('picking_type_id.name', 'in', ['Rackeo', 'Rackeos', 'Storage'])
        ])
        
        all_stors = (stor_pickings | rack_pickings).sorted('id')

        # 2. Recopilar move lines de los STOR
        move_lines = self.env['stock.move.line'].sudo().search([
            ('picking_id', 'in', all_stors.ids)
        ])

        prod_to_parent = {}
        po_products = {}

        for line in po.order_line.filtered(lambda l: not l.display_type and l.product_id):
            parent = line.product_id
            sku = parent.default_code or parent.name
            prod_to_parent[parent.id] = parent.id
            
            comp_prods = set([parent.id])
            prefix = (parent.default_code or '').strip()
            
            prefix_boxes = self.env['product.product'].sudo().search([('default_code', '=like', f"{prefix}#%")]) if prefix else self.env['product.product']
            if prefix_boxes:
                for box in prefix_boxes:
                    comp_prods.add(box.id)
                    prod_to_parent[box.id] = parent.id
            else:
                bom = self.env['mrp.bom'].sudo().search([
                    '|', ('product_id', '=', parent.id), ('product_tmpl_id', '=', parent.product_tmpl_id.id),
                    ('type', '=', 'phantom')
                ], limit=1)
                if bom and bom.bom_line_ids:
                    for bl in bom.bom_line_ids:
                        comp_prods.add(bl.product_id.id)
                        prod_to_parent[bl.product_id.id] = parent.id

            if parent.id not in po_products:
                po_products[parent.id] = {
                    'product': parent,
                    'sku': sku,
                    'ordered_qty': 0.0,
                    'received_qty': 0.0,
                    'racked_qty': 0.0,
                    'waiting_qty': 0.0,
                    'comp_ids': comp_prods,
                    'comp_racked_done': {cid: 0.0 for cid in comp_prods},
                    'comp_racked_waiting': {cid: 0.0 for cid in comp_prods},
                    'lines': []
                }
            po_products[parent.id]['ordered_qty'] += line.product_qty
            po_products[parent.id]['received_qty'] += line.qty_received

        processed_ml_pickings = set()

        for ml in move_lines:
            prod = ml.product_id
            if not prod:
                continue

            parent_id = prod_to_parent.get(prod.id)
            if not parent_id or parent_id not in po_products:
                sku = prod.default_code or prod.name
                po_products[prod.id] = {
                    'product': prod,
                    'sku': sku,
                    'ordered_qty': 0.0,
                    'received_qty': 0.0,
                    'racked_qty': 0.0,
                    'waiting_qty': 0.0,
                    'comp_ids': set([prod.id]),
                    'comp_racked_done': {prod.id: 0.0},
                    'comp_racked_waiting': {prod.id: 0.0},
                    'lines': []
                }
                parent_id = prod.id
                prod_to_parent[prod.id] = prod.id

            pdata = po_products[parent_id]
            picking = ml.picking_id
            processed_ml_pickings.add(picking.id)
            user_name = picking.operator.name or picking.create_uid.name or 'Sistema'
            is_automated = picking.validated_by_automation or ('automatizad' in (picking.wmds_log.mapped('log') and ' '.join(picking.wmds_log.mapped('log')).lower() or ''))
            is_done = (picking.state == 'done')
            is_waiting = (picking.state not in ('done', 'cancel'))

            qty = ml.quantity or ml.qty_done or ml.product_uom_qty or 0.0
            move_date = ml.write_date or ml.create_date or picking.date_done or picking.create_date
            date_str = move_date.strftime('%Y-%m-%d %H:%M:%S') if move_date else ''

            if prod.id in pdata['comp_racked_done']:
                if is_done:
                    pdata['comp_racked_done'][prod.id] += qty
                elif is_waiting:
                    pdata['comp_racked_waiting'][prod.id] += qty

            line_data = {
                'po': po.name,
                'parent_sku': pdata['sku'],
                'sku': prod.default_code or prod.name,
                'product_name': prod.name,
                'origin_location': ml.location_id.complete_name or ml.location_id.name,
                'dest_location': ml.location_dest_id.complete_name or ml.location_dest_id.name,
                'quantity': qty,
                'ok': '✅' if is_done else ('⏳ En Espera' if is_waiting else '❌ Cancelado'),
                'stor': picking.name,
                'state': 'Realizado' if is_done else ('En Espera' if is_waiting else picking.state),
                'user': user_name,
                'is_automated': 'SI' if is_automated else 'NO',
                'date': date_str,
                'is_done': is_done,
                'is_waiting': is_waiting,
            }

            pdata['lines'].append(line_data)

        # Compute kit equivalent racked and waiting quantities
        for parent_id, pdata in po_products.items():
            box_cids = [cid for cid in pdata['comp_ids'] if cid != parent_id]
            if box_cids:
                pdata['racked_qty'] = min(pdata['comp_racked_done'][cid] for cid in box_cids)
                pdata['waiting_qty'] = min(pdata['comp_racked_waiting'][cid] for cid in box_cids)
            else:
                pdata['racked_qty'] = pdata['comp_racked_done'].get(parent_id, 0.0)
                pdata['waiting_qty'] = pdata['comp_racked_waiting'].get(parent_id, 0.0)

        # 3. Generar Libro Excel
        wb = openpyxl.Workbook()
        default_sheet = wb.active
        wb.remove(default_sheet)

        # Estilos Excel
        header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='1F497D', end_color='1F497D', fill_type='solid')
        sub_font = Font(name='Calibri', size=11, bold=True)
        over_demand_font = Font(name='Calibri', size=11, bold=True, color='C00000')
        waiting_font = Font(name='Calibri', size=11, bold=True, color='E67E22')
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        align_center = Alignment(horizontal='center', vertical='center')
        align_left = Alignment(horizontal='left', vertical='center')
        align_right = Alignment(horizontal='right', vertical='center')

        # --- HOJA 1: RESUMEN GENERAL ---
        ws_summary = wb.create_sheet(title='Resumen_General')
        ws_summary.append(['REPORTE DE RACKEO DE ORDEN DE COMPRA'])
        ws_summary.append([f'PO: {po.name}', f'Proveedor: {po.partner_id.name or ""}', f'Estado PO: {po.state}'])
        ws_summary.append([])

        summary_headers = [
            'SKU', 'Producto', 'Cant. Ordenada (PO)', 'Cant. Recibida (IN)',
            'Dif. Recepción vs PO', 'Estado Recepción',
            'Rackeado (Realizado)', 'En Espera (Pendiente)',
            'Faltantes (vs PO)', 'Sobredemanda (vs PO)',
            'Estado Rackeo (vs PO)', 'Estado Rackeo (vs Recepción)'
        ]
        ws_summary.append(summary_headers)

        for col_num in range(1, len(summary_headers) + 1):
            cell = ws_summary.cell(row=4, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center

        total_ordered = 0.0
        total_received = 0.0
        total_racked = 0.0
        total_waiting = 0.0
        total_overdemand_po = 0.0
        total_missing_po = 0.0

        curr_row = 5
        for prod_id, pdata in po_products.items():
            racked = pdata['racked_qty']
            waiting = pdata['waiting_qty']
            ordered = pdata['ordered_qty']
            received = pdata['received_qty']
            total_processed = racked + waiting

            # 1. Auditoría Recepción vs PO
            diff_rec = received - ordered
            if received > ordered and ordered > 0:
                status_rec = f"⚠️ Exceso Recibido (+{int(diff_rec) if diff_rec.is_integer() else diff_rec} pzs vs PO)"
            elif received < ordered and ordered > 0:
                missing_rec = ordered - received
                status_rec = f"⚠️ Recepción Incompleta (-{int(missing_rec) if missing_rec.is_integer() else missing_rec} pzs vs PO)"
            else:
                status_rec = "✅ Recepción Completa"

            # 2. Auditoría Rackeo vs PO (Base Target = Ordered)
            base_target_po = ordered if ordered > 0 else received
            missing_po = max(0.0, base_target_po - racked)
            overdemand_po = max(0.0, total_processed - base_target_po)

            if overdemand_po > 0:
                status_rack_po = f"sobre demanda de rackeo en {int(overdemand_po) if overdemand_po.is_integer() else overdemand_po}, borrar"
            elif missing_po > 0:
                if waiting > 0:
                    status_rack_po = f"⚠️ Incompleto / Faltan {int(missing_po) if missing_po.is_integer() else missing_po} pzs ({int(waiting) if waiting.is_integer() else waiting} en espera)"
                else:
                    status_rack_po = f"⚠️ Incompleto / Faltan {int(missing_po) if missing_po.is_integer() else missing_po} pzs"
            else:
                status_rack_po = '✅ Acomodado Completo'

            # 3. Auditoría Rackeo vs Recepción (Base Target = Received)
            base_target_rec = received if received > 0 else ordered
            overdemand_rec = max(0.0, total_processed - base_target_rec)
            missing_rec_rack = max(0.0, base_target_rec - racked)

            if overdemand_rec > 0:
                status_rack_rec = f"Exceso STOR vs Recepción ({int(overdemand_rec)} pzs)"
            elif missing_rec_rack > 0:
                status_rack_rec = f"Pendiente de Acomodar ({int(missing_rec_rack)} pzs por rackear)"
            else:
                status_rack_rec = "✅ Acomodado Completo (100% de lo recibido)"

            ws_summary.append([
                pdata['sku'],
                pdata['product'].name,
                ordered,
                received,
                diff_rec,
                status_rec,
                racked,
                waiting,
                missing_po,
                overdemand_po,
                status_rack_po,
                status_rack_rec
            ])

            ws_summary.cell(row=curr_row, column=1).alignment = align_left
            ws_summary.cell(row=curr_row, column=2).alignment = align_left
            ws_summary.cell(row=curr_row, column=3).alignment = align_right
            ws_summary.cell(row=curr_row, column=4).alignment = align_right
            ws_summary.cell(row=curr_row, column=5).alignment = align_right
            ws_summary.cell(row=curr_row, column=6).alignment = align_center
            ws_summary.cell(row=curr_row, column=7).alignment = align_right
            ws_summary.cell(row=curr_row, column=8).alignment = align_right
            ws_summary.cell(row=curr_row, column=9).alignment = align_right
            ws_summary.cell(row=curr_row, column=10).alignment = align_right
            ws_summary.cell(row=curr_row, column=11).alignment = align_center
            ws_summary.cell(row=curr_row, column=12).alignment = align_center

            if "borrar" in status_rack_po:
                ws_summary.cell(row=curr_row, column=11).font = over_demand_font
            elif "Incompleto" in status_rack_po:
                ws_summary.cell(row=curr_row, column=11).font = waiting_font

            if "Exceso" in status_rec:
                ws_summary.cell(row=curr_row, column=6).font = waiting_font

            for c in range(1, 13):
                ws_summary.cell(row=curr_row, column=c).border = thin_border

            total_ordered += ordered
            total_received += received
            total_racked += racked
            total_waiting += waiting
            total_missing_po += missing_po
            total_overdemand_po += overdemand_po
            curr_row += 1

        # Totales
        tot_diff_rec = total_received - total_ordered
        ws_summary.append([
            'TOTALES',
            '',
            total_ordered,
            total_received,
            tot_diff_rec,
            '⚠️ EXCESOS EN RECEPCIÓN' if tot_diff_rec > 0 else ('⚠️ FALTANTES EN RECEPCIÓN' if tot_diff_rec < 0 else '✅ RECEPCIÓN PERFECTA'),
            total_racked,
            total_waiting,
            total_missing_po,
            total_overdemand_po,
            '⚠️ REVISAR SOBREDEMANDAS DE RACKEO' if total_overdemand_po > 0 else ('❌ PO CON PENDIENTES DE RACKEO' if total_missing_po > 0 else '✅ Acomodado Completo'),
            '✅ Acomodado Completo vs Recepción' if (total_racked + total_waiting >= total_received and total_received > 0) else '⚠️ PENDIENTES DE RACKEO VS RECEPCIÓN'
        ])

        for c in range(1, 13):
            cell = ws_summary.cell(row=curr_row, column=c)
            cell.font = sub_font
            cell.border = thin_border

        # --- HOJAS POR PRODUCTO (SKU) ---
        used_sheet_names = set(['Resumen_General'])

        for prod_id, pdata in po_products.items():
            raw_sku = pdata['sku'] or 'SKU_DESCONOCIDO'
            clean_sku = "".join([c for c in raw_sku if c.isalnum() or c in ('-', '_')]).strip()[:28]
            sheet_title = clean_sku or f'Prod_{prod_id}'

            base_title = sheet_title
            counter = 1
            while sheet_title in used_sheet_names:
                sheet_title = f"{base_title[:25]}_{counter}"
                counter += 1
            used_sheet_names.add(sheet_title)

            ws_prod = wb.create_sheet(title=sheet_title)

            racked = pdata['racked_qty']
            waiting = pdata['waiting_qty']
            ordered = pdata['ordered_qty']
            received = pdata['received_qty']
            total_proc = racked + waiting
            base_target = ordered if ordered > 0 else received
            missing = max(0.0, base_target - racked)
            overdemand = max(0.0, total_proc - base_target)

            ws_prod.append([f"DESGLOSE DE MOVIMIENTOS RACKEO (STOR) - SKU: {pdata['sku']}"])
            ws_prod.append([f"Producto: {pdata['product'].name}"])
            
            if overdemand > 0:
                status_header_str = f"SOBREDEMANDA VS PO: sobre demanda de rackeo en {int(overdemand) if overdemand.is_integer() else overdemand}, borrar"
            elif missing > 0:
                if waiting > 0:
                    status_header_str = f"ESTADO VS PO: Faltan {int(missing) if missing.is_integer() else missing} pzs por rackear ({int(waiting) if waiting.is_integer() else waiting} en espera)"
                else:
                    status_header_str = f"ESTADO VS PO: Faltan {int(missing) if missing.is_integer() else missing} pzs por rackear"
            else:
                status_header_str = "ESTADO VS PO: ✅ Acomodado Completo"

            ws_prod.append([f"Ordenado PO: {ordered}", f"Recibido IN: {received}", f"Rackeado: {racked}", f"En Espera: {waiting}", status_header_str])
            ws_prod.append([])

            prod_headers = ['PO', 'SKU', 'FECHA Y HORA', 'ORIGEN', 'DESTINO (UBICACIÓN)', 'PZS', 'Ok / Estado', 'STOR', 'Estado Picking', 'Usuario', 'Automatizado', 'Observación']
            ws_prod.append(prod_headers)

            for col_num in range(1, len(prod_headers) + 1):
                cell = ws_prod.cell(row=5, column=col_num)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = align_center

            p_row = 6
            if pdata['lines']:
                accum_total = 0.0
                for line in pdata['lines']:
                    line_qty = line['quantity']
                    accum_total += line_qty

                    if accum_total > base_target and base_target > 0:
                        excess = min(line_qty, accum_total - base_target)
                        obs = f"sobre demanda de rackeo en {int(excess) if excess.is_integer() else excess}, borrar"
                    elif line.get('is_waiting'):
                        obs = "en espera"
                    else:
                        obs = "totales"

                    ws_prod.append([
                        line['po'],
                        line['sku'],
                        line['date'],
                        line['origin_location'],
                        line['dest_location'],
                        line['quantity'],
                        line['ok'],
                        line['stor'],
                        line['state'],
                        line['user'],
                        line['is_automated'],
                        obs
                    ])

                    ws_prod.cell(row=p_row, column=1).alignment = align_left
                    ws_prod.cell(row=p_row, column=2).alignment = align_left
                    ws_prod.cell(row=p_row, column=3).alignment = align_center
                    ws_prod.cell(row=p_row, column=4).alignment = align_left
                    ws_prod.cell(row=p_row, column=5).alignment = align_left
                    ws_prod.cell(row=p_row, column=6).alignment = align_right
                    ws_prod.cell(row=p_row, column=7).alignment = align_center
                    ws_prod.cell(row=p_row, column=8).alignment = align_left
                    ws_prod.cell(row=p_row, column=9).alignment = align_center
                    ws_prod.cell(row=p_row, column=10).alignment = align_left
                    ws_prod.cell(row=p_row, column=11).alignment = align_center
                    ws_prod.cell(row=p_row, column=12).alignment = align_left

                    if "borrar" in obs:
                        ws_prod.cell(row=p_row, column=12).font = over_demand_font
                    elif line.get('is_waiting'):
                        ws_prod.cell(row=p_row, column=7).font = waiting_font

                    for c in range(1, 13):
                        ws_prod.cell(row=p_row, column=c).border = thin_border
                    p_row += 1
            else:
                ws_prod.append(['- Sin movimientos de rackeo registrados -'])

        # Auto-ajustar ancho de columnas en todas las hojas
        for ws in wb.worksheets:
            for col in ws.columns:
                max_len = 0
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or '')
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 65)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        file_data = base64.b64encode(output.read())
        output.close()

        file_name = f"Reporte_Rackeo_{po.name}.xlsx"

        # Eliminar adjuntos anteriores de reportes generados previamente por este wizard para evitar acumulacion
        old_attachments = self.env['ir.attachment'].sudo().search([
            ('res_model', '=', self._name),
            ('res_id', '=', self.id)
        ])
        old_attachments.unlink()

        self.write({
            'report_file': file_data,
            'file_name': file_name,
        })

        attachment = self.env['ir.attachment'].sudo().create({
            'name': file_name,
            'datas': file_data,
            'res_model': self._name,
            'res_id': self.id,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        })

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }
